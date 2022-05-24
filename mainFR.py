from googletrans import Translator
from docx import Document
from documents import create_psd_fr, create_pfr_fr, create_test_plan_fr
from mailer import send_mail
from questions import get_id_fr
from todo import show_todo, show_history, insert_todo, finish_todo
from openpyxl import load_workbook


def showtodo(mail):
    todo = ""
    df = show_todo(mail)
    df = df.reset_index()
    for index, row in df.iterrows():
        todo = todo + str(index) + " - " + row['todo'] + ";"
    return todo


def showhistory(mail):
    todo = ""
    df = show_history(mail)
    df = df.reset_index()
    for index, row in df.iterrows():
        todo = todo + row['todo'] + " \"" + row['status'] + "\" ;"
    return todo


def translateToEnglish(text):
    translator = Translator()
    translation = translator.translate(text, dest="en")
    return translation.text


def email(msg, question, mail):
    send_mail(msg, mail, "Mail générer")
    return "Le courrier généré a été envoyé à votre adresse"


def add_to_do(msg, question, mail):
    if question == "what is the to do to add?":
        insert_todo(msg, mail)
        return "To do added successfully."


def complete_to_do(msg, question):
    if question == "what is the number of to do to complete?":
        try:
            finish_todo(msg)
        except ValueError:
            print("Invalid Number entered")
            return "number entered not valid"
        return "To do completed successfully."


def generate_answer_fr(msg, mail, question):
    question_id = get_id_fr(question)
    if question_id == 1:
        return create_psd_fr(msg, question, mail)
    elif question_id == 2:
        return create_pfr_fr(msg, question, mail)
    elif question_id == 3:
        return create_test_plan_fr(msg, question, mail)
    elif question_id == 4:
        return email(msg, question, mail)
    elif question_id == 5:
        return translateToEnglish(msg)
    elif question_id == 6:
        return add_to_do(msg, question, mail)
    elif question_id == 7:
        return complete_to_do(msg, question)
    else:
        return bot_functions_fr(msg.lower(), mail)


#     lang = 'fr'
#     st.sidebar.markdown("# Commande")
#     st.sidebar.markdown("Les commandes de l\'assistant vocale : ")
#     st.sidebar.markdown("1 - Vermera : pour utiliser le microphone")
#     st.sidebar.markdown("2 - mes taches : pour voir vos taches")
#     st.sidebar.markdown("3 - ajoute tache : pour ajouter une nouvelle tache")
#     st.sidebar.markdown("4 - compléter tache : pour completer une tache")
#     st.sidebar.markdown("5 - historique tache : pour voir l\'historique des taches")
#     st.sidebar.markdown("6 - traduire : traduire du français vers l'anglais")
#     st.sidebar.markdown("7 - document développement : créer PSD.")
#     st.sidebar.markdown("8 - document client : créer PFR.")
#     st.sidebar.markdown("9 - plan de test : créer un plan de test.")


def bot_functions_fr(query, mail):
    if check_keyword(query, ['mail', 'email', 'courriel']):
        return "Que devrais-je écrire?"
    elif check_keyword(query, ['traduire', 'anglais']):
        return "tu veux traduire quoi ?"

    elif check_keyword(query, ['mes tâches', 'mes taches']):
        return showtodo(mail)

    elif check_keyword(query, ['ajoute tache', 'ajoute tâche', 'nouvelle tache', 'nouvelle tâche']):
        return "quelle est la tâche à ajouter?"

    elif check_keyword(query, ['completer tache', 'compléter tache', 'compléter tâche']):
        return "quel est le numéro de la tâche à accomplir?"

    elif check_keyword(query, ['historique tâche', 'historique tache']):
        return showhistory(mail)
    elif check_keyword(query, ['psd', 'document dev', 'document de développement']):
        document = Document('utils/PSD.docx')
        document.save(mail + "PSD.docx")
        return "PSD en création, tu veux quoi comme titre ?"
    elif check_keyword(query, ['pfr', 'document client']):
        document = Document('utils/PFR.docx')
        document.save(mail + "PFR.docx")
        return "PFR en création, tu veux quoi comme titre ?"
    elif check_keyword(query, ['plan de test', 'test']):
        workbook = load_workbook(filename='utils/plan.xlsx')
        workbook.save(mail + "plan.xlsx")
        return "plan de test en création, quel est le titre du test ?"
    else:
        return "model_fr"


def check_keyword(query, keywords):
    for key in keywords:
        if key in query:
            return True
    return False
